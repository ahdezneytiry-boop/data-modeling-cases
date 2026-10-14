# -*- coding: utf-8 -*-
"""
从 Q4 全年 JSON 生成官方 result4-2.xlsx / result4-3.xlsx（复制附件5模板填写）
并回读核验：维度、无空值、与审计数字一致。
"""
import json, os, shutil
import numpy as np
import pandas as pd
import openpyxl

BASE = r'C:\Users\ahdezneytiry-boop1\.zcode\workspace\default\q4run'
TPL = os.path.join(BASE, '附件5')
OUT = os.path.join(BASE, '交付')
os.makedirs(OUT, exist_ok=True)
T = 144


def interval_label(t):
    """t=1..144 -> 'HH:MM-HH:MM' (右端点口径)"""
    s = (t - 1) * 10
    e = t * 10
    def fmt(m):
        return ('%d:%02d' % (m // 60, m % 60)) if m % 60 else ('24:00' if m == 1440 else '%d:00' % (m // 60))
    return '%s-%s' % (fmt(s), fmt(e))


def load_rec(run, tag):
    rec = json.load(open(os.path.join(run, tag + '.json'), encoding='utf-8'))
    ev = [r for r in rec if '2025-02-01' <= r['date'] <= '2025-12-31']
    assert len(ev) == 334, 'evaluation days != 334'
    return ev


def fill_plan(ws, rows, key):
    for i, r in enumerate(rows):
        vals = np.array(r[key])
        for j in range(T):
            ws.cell(row=2 + i, column=2 + j, value=round(float(vals[j]), 6))


def fill_storage(ws, rows):
    for i, r in enumerate(rows):
        c = np.array(r['c']); rr = np.array(r['r'])
        for b in range(6):
            row = 2 + i * 6 + b
            ws.cell(row=row, column=3, value=round(float(c[b * 24:(b + 1) * 24].sum()), 6))
            ws.cell(row=row, column=4, value=round(float(rr[b * 24:(b + 1) * 24].sum()), 6))
        ws.cell(row=2 + i * 6, column=5, value='0:00')
        ws.cell(row=2 + i * 6, column=6, value=round(float(r['E_in']), 6))
        ws.cell(row=2 + i * 6 + 1, column=5, value='24:00')
        ws.cell(row=2 + i * 6 + 1, column=6, value=round(float(r['E_end']), 6))


def fill_emergency(ws, rows):
    row = 2
    for r in rows:
        g = np.array(r['g'])
        pos = np.where(g > 1e-9)[0]
        if len(pos) == 0:
            ws.cell(row=row, column=1, value=pd.Timestamp(r['date']))
            ws.cell(row=row, column=3, value=0)
            row += 1
            continue
        # 合并相邻正 g 段
        segs = []
        start = pos[0]; prev = pos[0]
        for t in pos[1:]:
            if t == prev + 1:
                prev = t
            else:
                segs.append((start, prev)); start = t; prev = t
        segs.append((start, prev))
        for s, e in segs:
            ws.cell(row=row, column=1, value=pd.Timestamp(r['date']))
            lab = interval_label(s + 1).split('-')[0] + '-' + interval_label(e + 1).split('-')[1]
            ws.cell(row=row, column=2, value=lab)
            ws.cell(row=row, column=3, value=round(float(g[s:e + 1].sum()), 6))
            row += 1


def export(run_dir, tag, tpl_name, out_name, extra_sheet=None):
    shutil.copy(os.path.join(TPL, tpl_name), os.path.join(OUT, out_name))
    rec = load_rec(run_dir, tag)
    wb = openpyxl.load_workbook(os.path.join(OUT, out_name))
    ws = wb['计划购电量']
    for j in range(T):
        ws.cell(row=1, column=2 + j, value=interval_label(j + 1))
    fill_plan(ws, rec, 'q0')
    ws2 = wb['充放电量']
    fill_storage(ws2, rec)
    ws3 = wb['紧急购电量']
    fill_emergency(ws3, rec)
    if extra_sheet:
        ws4 = wb['调整购电量']
        for j in range(T):
            ws4.cell(row=1, column=2 + j, value=interval_label(j + 1))
        fill_plan(ws4, rec, 'qfinal')
    wb.save(os.path.join(OUT, out_name))
    print('exported', out_name)


export(os.path.join(BASE, '结果_refund'), 'Q42', 'result4-2.xlsx', 'result4-2.xlsx')
export(os.path.join(BASE, '结果_refund'), 'Q43', 'result4-3.xlsx', 'result4-3.xlsx', extra_sheet=True)

# ============ 回读核验 ============
print('\n===== 回读核验 =====')
errs = []
for fname, tag, has_adj in [('result4-2.xlsx', 'Q42', False), ('result4-3.xlsx', 'Q43', True)]:
    wb = openpyxl.load_workbook(os.path.join(OUT, fname))
    rec = load_rec(os.path.join(BASE, '结果_refund'), tag)
    ws = wb['计划购电量']
    if ws.max_row != 335 or ws.max_column != 147:
        errs.append('%s 计划购电量 dims %sx%s' % (fname, ws.max_row, ws.max_column))
    s = 0.0
    for i, r in enumerate(rec):
        vals = [ws.cell(row=2 + i, column=2 + j).value for j in range(T)]
        if any(v is None for v in vals):
            errs.append('%s 计划购电量 空值 row%d' % (fname, i))
        s += sum(vals)
    print('%s 计划购电量 回读合计=%.6f vs JSON=%.6f (diff %.3g)'
          % (fname, s, sum(np.array(r['q0']).sum() for r in rec), abs(s - sum(np.array(r['q0']).sum() for r in rec))))
    ws2 = wb['充放电量']
    csum = sum(ws2.cell(row=2 + i * 6 + b, column=3).value for i in range(334) for b in range(6))
    rsum = sum(ws2.cell(row=2 + i * 6 + b, column=4).value for i in range(334) for b in range(6))
    print('%s 充放电量 回读: 充电合计=%.4f 放电合计=%.4f (JSON %.4f / %.4f)'
          % (fname, csum, rsum, sum(np.array(r['c']).sum() for r in rec), sum(np.array(r['r']).sum() for r in rec)))
    ws3 = wb['紧急购电量']
    gsum = 0
    for row in ws3.iter_rows(min_row=2, max_col=3, values_only=True):
        if row[2] is not None:
            gsum += float(row[2])
    print('%s 紧急购电量 回读合计=%.4f (JSON %.4f)' % (fname, gsum, sum(np.array(r['g']).sum() for r in rec)))
    if has_adj:
        ws4 = wb['调整购电量']
        asum = 0
        for i, r in enumerate(rec):
            vals = [ws4.cell(row=2 + i, column=2 + j).value for j in range(T)]
            if any(v is None for v in vals):
                errs.append('%s 调整购电量 空值 row%d' % (fname, i))
            asum += sum(vals)
        print('%s 调整购电量 回读合计=%.6f vs JSON=%.6f (diff %.3g)'
              % (fname, asum, sum(np.array(r['qfinal']).sum() for r in rec),
                 abs(asum - sum(np.array(r['qfinal']).sum() for r in rec))))
print('回读错误数:', len(errs))
for e in errs:
    print('  ERR:', e)
print('EXPORT', 'PASS' if not errs else 'FAIL')
